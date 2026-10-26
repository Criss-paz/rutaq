#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generar_presentacion.py

Genera la presentación formal en PowerPoint (16:9) del Proyecto II de
Investigación de Operaciones — Caso Transporte Express (RutaQ).
Incluye 12 diapositivas con diseño moderno, tipografía legible, capturas
reales del sistema y notas completas del orador tomadas del guion.

Salida:
    docs/proyecto2/presentacion/Presentacion_Transporte_Express.pptx
"""

import os
import sys
from pathlib import Path
from PIL import Image

import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# -----------------------------------------------------------------------------
# Constantes de Diseño y Paleta de Colores (RutaQ)
# -----------------------------------------------------------------------------
SLIDE_WIDTH_IN = 13.333
SLIDE_HEIGHT_IN = 7.5

FONT_FAMILY = "Segoe UI"
FONT_MONO = "Consolas"

# Paleta basada en la interfaz oficial de RutaQ
BG_DARK = RGBColor(11, 15, 25)         # Fondo principal slate profundo (#0B0F19)
BG_CARD = RGBColor(24, 32, 47)         # Superficie de tarjetas (#18202F)
BG_CARD_ALT = RGBColor(18, 25, 38)     # Superficie secundaria (#121926)
BORDER_CARD = RGBColor(51, 65, 85)     # Borde sutil slate (#334155)
BORDER_ACCENT = RGBColor(71, 85, 105)  # Borde destacado (#475569)

# Acentos oficiales
ACCENT_GREEN = RGBColor(16, 185, 129)   # Verde Transporte Express (#10B981)
ACCENT_GREEN_DARK = RGBColor(5, 150, 105) # #059669
ACCENT_BLUE = RGBColor(59, 130, 246)    # Azul interactivo (#3B82F6)
ACCENT_PURPLE = RGBColor(168, 85, 247)  # Morado de ruta óptima (#A855F7)
ACCENT_AMBER = RGBColor(245, 158, 11)   # Ámbar alerta/resultado (#F59E0B)
ACCENT_CYAN = RGBColor(6, 182, 212)     # Cyan complementario (#06B6D4)

# Colores de texto
TEXT_WHITE = RGBColor(248, 250, 252)    # Texto primario blanco (#F8FAFC)
TEXT_MUTED = RGBColor(148, 163, 184)    # Texto secundario (#94A3B8)
TEXT_LIGHT = RGBColor(226, 232, 240)    # Texto cuerpo destacado (#E2E8F0)

# Fondos de insignia (badges)
BADGE_BG_GREEN = RGBColor(6, 78, 59)
BADGE_BG_BLUE = RGBColor(30, 58, 138)
BADGE_BG_PURPLE = RGBColor(88, 28, 135)

# Directorios de referencia
BASE_DIR = Path(__file__).resolve().parent.parent
CAPTURES_DIR = BASE_DIR / "docs" / "capturas" / "proyecto2"
CAPTURES_ENSAYO_DIR = CAPTURES_DIR / "ensayo"
OUTPUT_DIR = BASE_DIR / "docs" / "proyecto2" / "presentacion"
OUTPUT_FILE = OUTPUT_DIR / "Presentacion_Transporte_Express.pptx"


def ensure_dirs():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def create_blank_slide(prs):
    """Crea una diapositiva en blanco con fondo oscuro uniforme."""
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Fondo oscuro completo
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(SLIDE_WIDTH_IN), Inches(SLIDE_HEIGHT_IN)
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_DARK
    bg.line.fill.background()
    return slide


def add_badge(slide, text, left, top, width=Inches(3.2), height=Inches(0.32),
              fg_color=ACCENT_GREEN, bg_color=BADGE_BG_GREEN):
    """Agrega una insignia estilizada con texto en mayúsculas."""
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    badge.fill.solid()
    badge.fill.fore_color.rgb = bg_color
    badge.line.color.rgb = fg_color
    badge.line.width = Pt(1.0)

    tf = badge.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.0)
    tf.margin_bottom = Inches(0.0)

    p = tf.paragraphs[0]
    p.text = text.upper()
    p.alignment = PP_ALIGN.LEFT
    p.font.name = FONT_FAMILY
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = fg_color
    return badge


def add_header(slide, badge_text, title_text, subtitle_text=None,
               badge_fg=ACCENT_GREEN, badge_bg=BADGE_BG_GREEN):
    """Agrega el encabezado superior estándar de la diapositiva."""
    left = Inches(0.8)
    # Insignia
    add_badge(slide, badge_text, left, Inches(0.42), width=Inches(3.4), height=Inches(0.30),
              fg_color=badge_fg, bg_color=badge_bg)

    # Título principal
    tx_box = slide.shapes.add_textbox(left, Inches(0.76), Inches(11.7), Inches(0.65))
    tf = tx_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = FONT_FAMILY
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    # Subtítulo opcional
    if subtitle_text:
        sub_box = slide.shapes.add_textbox(left, Inches(1.38), Inches(11.7), Inches(0.40))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle_text
        p_sub.font.name = FONT_FAMILY
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = TEXT_MUTED


def add_card(slide, left, top, width, height, bg_color=BG_CARD, border_color=BORDER_CARD):
    """Crea una tarjeta contenedora estilizada."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    return card


def add_metric_box(slide, left, top, width, height, value_str, label_str,
                   accent_color=ACCENT_GREEN, bg_color=BG_CARD):
    """Crea un bloque de métrica destacada con número grande y etiqueta descriptiva."""
    card = add_card(slide, left, top, width, height, bg_color=bg_color, border_color=accent_color)
    tx = slide.shapes.add_textbox(left + Inches(0.12), top + Inches(0.08), width - Inches(0.24), height - Inches(0.16))
    tf = tx.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p_val = tf.paragraphs[0]
    p_val.text = value_str
    p_val.font.name = FONT_FAMILY
    p_val.font.size = Pt(24)
    p_val.font.bold = True
    p_val.font.color.rgb = accent_color
    p_val.alignment = PP_ALIGN.CENTER

    p_lbl = tf.add_paragraph()
    p_lbl.text = label_str
    p_lbl.font.name = FONT_FAMILY
    p_lbl.font.size = Pt(10.5)
    p_lbl.font.bold = True
    p_lbl.font.color.rgb = TEXT_LIGHT
    p_lbl.alignment = PP_ALIGN.CENTER
    return card


def add_image_fitted(slide, img_path, box_left, box_top, box_w, box_h,
                     caption_text=None, border_color=BORDER_CARD):
    """Inserta una imagen proporcionalmente dentro del rectángulo dado con marco estilizado."""
    img_path = Path(img_path)
    if not img_path.exists():
        print(f"[AVISO] Imagen no encontrada: {img_path}")
        return

    with Image.open(img_path) as im:
        iw, ih = im.size

    aspect = iw / ih
    box_aspect = (box_w.inches) / (box_h.inches)

    if aspect > box_aspect:
        w_in = box_w.inches
        h_in = w_in / aspect
    else:
        h_in = box_h.inches
        w_in = h_in * aspect

    final_w = Inches(w_in)
    final_h = Inches(h_in)
    final_left = box_left + (box_w - final_w) / 2
    final_top = box_top + (box_h - final_h) / 2

    # Marco posterior
    frame = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        final_left - Inches(0.04), final_top - Inches(0.04),
        final_w + Inches(0.08), final_h + Inches(0.08)
    )
    frame.fill.solid()
    frame.fill.fore_color.rgb = BG_CARD
    frame.line.color.rgb = border_color
    frame.line.width = Pt(1.5)

    # Imagen
    slide.shapes.add_picture(str(img_path), final_left, final_top, width=final_w, height=final_h)

    # Pie de foto opcional
    if caption_text:
        cap_box = slide.shapes.add_textbox(
            box_left, final_top + final_h + Inches(0.06), box_w, Inches(0.32)
        )
        tf_cap = cap_box.text_frame
        tf_cap.word_wrap = True
        tf_cap.margin_left = tf_cap.margin_right = tf_cap.margin_top = tf_cap.margin_bottom = 0
        p_cap = tf_cap.paragraphs[0]
        p_cap.text = caption_text
        p_cap.font.name = FONT_FAMILY
        p_cap.font.size = Pt(9.5)
        p_cap.font.color.rgb = TEXT_MUTED
        p_cap.alignment = PP_ALIGN.CENTER


def set_speaker_notes(slide, notes_text):
    """Establece las notas del orador en la diapositiva."""
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes_text.strip()


# -----------------------------------------------------------------------------
# Constructores de cada Diapositiva
# -----------------------------------------------------------------------------

def build_slide_01(prs):
    """Diapositiva 1: Portada."""
    slide = create_blank_slide(prs)

    # Tarjeta de Portada Principal
    card_w = Inches(11.733)
    card_h = Inches(6.3)
    card_left = Inches(0.8)
    card_top = Inches(0.6)
    card = add_card(slide, card_left, card_top, card_w, card_h, bg_color=BG_CARD, border_color=BORDER_ACCENT)

    # Insignia Institucional
    add_badge(slide, "UNIVERSIDAD MARIANO GÁLVEZ DE GUATEMALA · INGENIERÍA EN SISTEMAS",
              card_left + Inches(0.6), card_top + Inches(0.55),
              width=Inches(6.2), height=Inches(0.34),
              fg_color=ACCENT_BLUE, bg_color=BADGE_BG_BLUE)

    # Título principal
    title_box = slide.shapes.add_textbox(card_left + Inches(0.6), card_top + Inches(1.05), Inches(10.5), Inches(1.3))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p1 = tf.paragraphs[0]
    p1.text = "Transporte Express — Ruta Más Corta"
    p1.font.name = FONT_FAMILY
    p1.font.size = Pt(34)
    p1.font.bold = True
    p1.font.color.rgb = ACCENT_GREEN

    p2 = tf.add_paragraph()
    p2.text = "Optimización de Rutas Departamentales y Red Troncal Mínima (RutaQ)"
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(17)
    p2.font.color.rgb = TEXT_LIGHT

    # Detalles del Curso y Catedrático
    info_box = slide.shapes.add_textbox(card_left + Inches(0.6), card_top + Inches(2.45), Inches(5.2), Inches(2.8))
    tf_info = info_box.text_frame
    tf_info.word_wrap = True
    tf_info.margin_left = tf_info.margin_top = tf_info.margin_right = tf_info.margin_bottom = 0

    p = tf_info.paragraphs[0]
    p.text = "DATOS ACADÉMICOS"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    items_info = [
        ("Curso:", "Investigación de Operaciones"),
        ("Catedrático:", "Ing. Samuel Isaac Gómez Moreno"),
        ("Semestre:", "6.º semestre, Ingeniería en Sistemas"),
        ("Fecha:", "Sábado 31 de octubre de 2026"),
        ("Plataforma:", "https://criss-paz.github.io/rutaq/"),
    ]
    for lbl, val in items_info:
        p = tf_info.add_paragraph()
        p.text = f"{lbl} {val}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE

    # Tarjeta de Integrantes
    int_box = slide.shapes.add_textbox(card_left + Inches(6.0), card_top + Inches(2.45), Inches(5.1), Inches(2.8))
    tf_int = int_box.text_frame
    tf_int.word_wrap = True
    tf_int.margin_left = tf_int.margin_top = tf_int.margin_right = tf_int.margin_bottom = 0

    p = tf_int.paragraphs[0]
    p.text = "INTEGRANTES DEL EQUIPO"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_PURPLE

    integrantes = [
        ("Cristopher Alexis Castellanos Paz", "Carné: 2690-24-5972"),
        ("Leonel Taracena Marroquín", "Carné: 2690-24-4072"),
        ("Pablo Alexander Vásquez Rivera", "Carné: 2690-24-13123"),
    ]
    for nombre, carne in integrantes:
        p_n = tf_int.add_paragraph()
        p_n.text = f"• {nombre}"
        p_n.font.name = FONT_FAMILY
        p_n.font.size = Pt(13.5)
        p_n.font.bold = True
        p_n.font.color.rgb = TEXT_WHITE

        p_c = tf_int.add_paragraph()
        p_c.text = f"   {carne}"
        p_c.font.name = FONT_FAMILY
        p_c.font.size = Pt(11.5)
        p_c.font.color.rgb = TEXT_MUTED

    # Pie de portada
    footer_box = slide.shapes.add_textbox(card_left + Inches(0.6), card_top + Inches(5.55), Inches(10.5), Inches(0.4))
    tf_f = footer_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "22 Cabeceras Departamentales · 49 Carreteras Directas · Algoritmos: Dijkstra, Prim y Kruskal"
    p_f.font.name = FONT_FAMILY
    p_f.font.size = Pt(11)
    p_f.font.bold = True
    p_f.font.color.rgb = ACCENT_AMBER
    p_f.alignment = PP_ALIGN.CENTER

    set_speaker_notes(
        slide,
        "Buenos días ingeniero Samuel Gómez y compañeros. Hoy les presentamos la solución al Proyecto II de "
        "Investigación de Operaciones: el Sistema de Optimización de Rutas Departamentales y Red Troncal Mínima "
        "para la empresa Transporte Express, desarrollado sobre nuestra plataforma web RutaQ. "
        "El equipo está integrado por Cristopher Alexis Castellanos Paz, Leonel Taracena Marroquín y "
        "Pablo Alexander Vásquez Rivera, estudiantes del 6.º semestre de Ingeniería en Sistemas de la "
        "Universidad Mariano Gálvez de Guatemala."
    )


def build_slide_02(prs):
    """Diapositiva 2: El Problema Logístico."""
    slide = create_blank_slide(prs)
    add_header(slide, "DIAGNÓSTICO Y REQUERIMIENTOS",
               "El Problema: Retrasos en Envíos Nacionales",
               "Causas operativas en Transporte Express y especificaciones del Proyecto II")

    # Columna izquierda: 2 tarjetas
    c1_left = Inches(0.8)
    c1_top = Inches(1.95)
    c1_w = Inches(5.6)
    c1_h = Inches(2.4)
    add_card(slide, c1_left, c1_top, c1_w, c1_h)

    tb1 = slide.shapes.add_textbox(c1_left + Inches(0.25), c1_top + Inches(0.2), c1_w - Inches(0.5), c1_h - Inches(0.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "PROBLEMÁTICA EN TRANSPORTE EXPRESS"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER

    points1 = [
        "Paquetes que llegan tarde sistemáticamente a su destino.",
        "Rutas elegidas por intuición o hábitos individuales sin sustento matemático.",
        "Desvíos innecesarios que provocan sobrecosto de combustible y desgaste de flota.",
        "Inexistencia de una red vial digitalizada y cuantificada en distancias reales."
    ]
    for pt in points1:
        p = tf1.add_paragraph()
        p.text = f"• {pt}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_LIGHT

    # Tarjeta 2: Requerimientos
    c2_top = Inches(4.55)
    c2_h = Inches(2.4)
    add_card(slide, c1_left, c2_top, c1_w, c2_h)

    tb2 = slide.shapes.add_textbox(c1_left + Inches(0.25), c2_top + Inches(0.2), c1_w - Inches(0.5), c2_h - Inches(0.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "REQUERIMIENTOS FORMALES DEL ENUNCIADO"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    points2 = [
        "1. Ingresar los 22 departamentos con cabeceras y coordenadas georreferenciadas.",
        "2. Registrar carreteras interdepartamentales con distancias reales en kilómetros.",
        "3. Preguntar interactivamente: ¿Dónde está el paquete? y ¿A dónde debe entregarse?",
        "4. Calcular y exhibir la mejor ruta con desglose de tramos y fórmula matemática."
    ]
    for pt in points2:
        p = tf2.add_paragraph()
        p.text = f"• {pt}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_LIGHT

    # Columna derecha: Imagen
    img_left = Inches(6.65)
    img_top = Inches(1.95)
    img_w = Inches(5.88)
    img_h = Inches(4.9)
    img_path = CAPTURES_ENSAYO_DIR / "01-pantalla-inicio-guatemala.png"
    add_image_fitted(slide, img_path, img_left, img_top, img_w, img_h,
                     caption_text="Figura 1: Vista general de RutaQ orientada al norte con los 22 departamentos")

    set_speaker_notes(
        slide,
        "Transporte Express es una compañía guatemalteca de logística que distribuye encomiendas a nivel nacional. "
        "La problemática central era que sus paquetes llegaban sistemáticamente tarde a destino debido a la falta "
        "de planificación. Los conductores tomaban rutas empíricas o hábitos individuales con desvíos innecesarios, "
        "sin contar con una red vial cuantificada con distancias reales ni una herramienta automatizada de despacho.\n\n"
        "El enunciado del Proyecto II solicita ingresar los 22 departamentos, registrar carreteras con kilometrajes reales, "
        "responder a las preguntas clave de dónde está el paquete y hacia dónde va, y presentar la ruta más corta "
        "de manera clara, visual y matemáticamente respaldada."
    )


def build_slide_03(prs):
    """Diapositiva 3: Modelo Matemático del Grafo."""
    slide = create_blank_slide(prs)
    add_header(slide, "MODELADO MATEMÁTICO",
               "Modelo: Grafo No Dirigido Ponderado G = (V, E, W)",
               "Abstracción determinista de la red vial nacional y criterios de conectividad")

    # Columna izquierda: 3 tarjetas métricas/conceptuales
    c_left = Inches(0.8)
    c_w = Inches(5.6)

    # Nodos
    add_card(slide, c_left, Inches(1.95), c_w, Inches(1.5))
    tb_n = slide.shapes.add_textbox(c_left + Inches(0.2), Inches(2.05), c_w - Inches(0.4), Inches(1.3))
    tf_n = tb_n.text_frame
    tf_n.word_wrap = True
    p = tf_n.paragraphs[0]
    p.text = "V = 22 Vértices (Cabeceras Departamentales)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    p = tf_n.add_paragraph()
    p.text = "• Coordenadas georreferenciadas oficiales (latitud/longitud) proyectadas en Web Mercator.\n" \
             "• Códigos canónicos inmutables de 3 letras (GUA, QUE, PET, IZA, SMA, etc.)."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # Aristas
    add_card(slide, c_left, Inches(3.60), c_w, Inches(1.75))
    tb_e = slide.shapes.add_textbox(c_left + Inches(0.2), Inches(3.70), c_w - Inches(0.4), Inches(1.55))
    tf_e = tb_e.text_frame
    tf_e.word_wrap = True
    p = tf_e.paragraphs[0]
    p.text = "E = 49 Aristas Directas entre Vecinos"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    p = tf_e.add_paragraph()
    p.text = "• Vecindad Territorial: Frontera geográfica compartida (contacto ≤ 200 m).\n" \
             "• Umbral de Paso: No cruzar a < 2.0 km de otra cabecera intermedia.\n" \
             "• Descarte Auditado: JAL–ZAC y PET–QUI descartadas formalmente por violar umbral."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # Pesos
    add_card(slide, c_left, Inches(5.50), c_w, Inches(1.45))
    tb_w = slide.shapes.add_textbox(c_left + Inches(0.2), Inches(5.60), c_w - Inches(0.4), Inches(1.25))
    tf_w = tb_w.text_frame
    tf_w.word_wrap = True
    p = tf_w.paragraphs[0]
    p.text = "W(e) Pesos: Kilometraje Real de Conducción"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_PURPLE

    p = tf_w.add_paragraph()
    p.text = "• Distancias físicas estrictamente positivas (W > 0).\n" \
             "• Nomenclatura vial oficial nacional (CA-1, CA-9, RN-15, FTN).\n" \
             "• Grafo 100 % conexo: existe ruta entre cualquier par de cabeceras."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # Columna derecha: Imagen del mapa completo
    img_left = Inches(6.65)
    img_top = Inches(1.95)
    img_w = Inches(5.88)
    img_h = Inches(4.9)
    img_path = CAPTURES_DIR / "02-mapa-completo.png"
    add_image_fitted(slide, img_path, img_left, img_top, img_w, img_h,
                     caption_text="Figura 2: Grafo completo: 22 cabeceras departamentales y 49 carreteras directas")

    set_speaker_notes(
        slide,
        "Para resolver el problema con rigor científico, modelamos la red vial de Guatemala como un grafo no dirigido "
        "ponderado y conexo G = (V, E, W). Contamos con 22 vértices que corresponden a las cabeceras departamentales "
        "oficiales con coordenadas georreferenciadas exactas en proyección Web Mercator con orientación norte-arriba.\n\n"
        "Para definir las carreteras directas aplicamos un criterio objetivo y reproducible: los departamentos deben "
        "compartir frontera territorial y la carretera no debe cruzar a menos de 2 km de otra cabecera intermedia. "
        "De este modo descartamos conexiones compuestas o atajos irreales (como JAL-ZAC y PET-QUI), obteniendo 49 aristas "
        "directas con pesos estrictamente positivos correspondientes a distancias reales de conducción en kilómetros."
    )


def build_slide_04(prs):
    """Diapositiva 4: Datos Reales y Validación."""
    slide = create_blank_slide(prs)
    add_header(slide, "PRECISIÓN Y METODOLOGÍA CARTOGRÁFICA",
               "Datos Reales y Auditoría Geográfica",
               "Extracción satelital con OSRM/OpenStreetMap y verificación al 100 % en Google Maps")

    c_left = Inches(0.8)
    c_w = Inches(5.6)

    # 3 Cajas Métricas Superiores
    metric_w = Inches(1.75)
    metric_h = Inches(1.2)
    gap = Inches(0.17)

    add_metric_box(slide, c_left, Inches(1.95), metric_w, metric_h,
                   "35 / 49", "Ruta Recomendada\n(OSRM a ±10 %)", ACCENT_GREEN)
    add_metric_box(slide, c_left + metric_w + gap, Inches(1.95), metric_w, metric_h,
                   "9 / 49", "Ruta Alternativa\n(OSRM a ±10 %)", ACCENT_BLUE)
    add_metric_box(slide, c_left + (metric_w + gap) * 2, Inches(1.95), metric_w, metric_h,
                   "5 / 49", "Valor Google Maps\n(Sin coincidencia)", ACCENT_AMBER)

    # Tarjeta de Metodología
    add_card(slide, c_left, Inches(3.35), c_w, Inches(3.60))
    tb_m = slide.shapes.add_textbox(c_left + Inches(0.25), Inches(3.50), c_w - Inches(0.5), Inches(3.30))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True

    p = tf_m.paragraphs[0]
    p.text = "METODOLOGÍA DE RECOLECCIÓN Y AUDITORÍA"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    bullets = [
        "Fuente Primaria: Geocodificación con Nominatim y cálculo de distancias de conducción sobre OpenStreetMap mediante OSRM.",
        "Auditoría 49/49 en Google Maps: Cada una de las 49 carreteras fue validada punto a punto en Google Maps.",
        "Tolerancia ±10 %: 35 tramos coinciden con la ruta recomendada y 9 con una ruta alternativa de Google Maps.",
        "5 Tramos Ajustados a Google Maps: Al no coincidir OSRM a ±10 %, se usa el valor de Google Maps (CHI-SOL, CHI-SUC, JAL-PRO, IZA-PET, QUE-TOT).",
        "Caso Crítico IZA–PET (287.0 km): Conexión directa por la CA-13 adoptada de Google Maps (OSRM daba 387.1 km) para evitar desvíos falsos por Alta Verapaz."
    ]
    for b in bullets:
        p = tf_m.add_paragraph()
        p.text = f"• {b}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_LIGHT

    # Columna derecha: Imagen del catálogo de carreteras
    img_left = Inches(6.65)
    img_top = Inches(1.95)
    img_w = Inches(5.88)
    img_h = Inches(4.9)
    img_path = CAPTURES_DIR / "04-formulario-carreteras.png"
    add_image_fitted(slide, img_path, img_left, img_top, img_w, img_h,
                     caption_text="Figura 3: Catálogo interactivo de las 49 carreteras auditadas con kilometraje real")

    set_speaker_notes(
        slide,
        "Un aspecto fundamental de nuestro trabajo fue garantizar que los datos fueran 100 % reales y auditables. "
        "Extrajimos las distancias y geometrías mediante el motor OSRM sobre OpenStreetMap y validamos manualmente "
        "las 49 carreteras frente a Google Maps bajo un criterio técnico de tolerancia a ±10 %:\n\n"
        "35 tramos con OSRM a ±10 % de la ruta recomendada por Google Maps; 9 tramos con OSRM a ±10 % de una ruta "
        "alternativa que ofrece Google Maps; y 5 tramos donde ninguna coincidió (CHI-SOL, CHI-SUC, JAL-PRO, IZA-PET, "
        "QUE-TOT), en los cuales se utiliza directamente el valor de Google Maps.\n\n"
        "Un caso crítico fue la arista Izabal–Petén de 287.0 km por la CA-13: OSRM arrojaba 387.1 km por un trazado "
        "indirecto, por lo que adoptar el valor oficial de Google Maps evita desviar las encomiendas por Alta Verapaz. "
        "Cada número en nuestro sistema proviene de carreteras reales de Guatemala debidamente auditadas."
    )


def build_slide_05(prs):
    """Diapositiva 5: Algoritmo de Dijkstra."""
    slide = create_blank_slide(prs)
    add_header(slide, "FUNDAMENTO ALGORÍTMICO",
               "Dijkstra: Algoritmo de Camino Mínimo",
               "Búsqueda voraz determinista para despacho punto a punto en la red vial")

    c_left = Inches(0.8)
    c_w = Inches(5.6)

    # Tarjeta 1: Principio y Relajación
    add_card(slide, c_left, Inches(1.95), c_w, Inches(2.25))
    tb_p = slide.shapes.add_textbox(c_left + Inches(0.25), Inches(2.05), c_w - Inches(0.5), Inches(2.05))
    tf_p = tb_p.text_frame
    tf_p.word_wrap = True
    p = tf_p.paragraphs[0]
    p.text = "PRINCIPIO VORAZ (GREEDY) Y RELAJACIÓN"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_PURPLE

    p = tf_p.add_paragraph()
    p.text = "• Enfoque: Encuentra el camino más corto desde un nodo origen s hacia todos los demás nodos de la red.\n" \
             "• Mecanismo de Relajación: Para cada arista (u, v) con peso W(u, v):\n" \
             "    Si  D[u] + W(u, v) < D[v]  entonces:\n" \
             "        D[v] ← D[u] + W(u, v)    y    prev[v] ← u\n" \
             "• Optimalidad: Al procesar un nodo u desde la cola de prioridad, su distancia D[u] es matemáticamente definitiva."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # Tarjeta 2: Complejidad y Pesos no negativos
    add_card(slide, c_left, Inches(4.35), c_w, Inches(2.60))
    tb_c = slide.shapes.add_textbox(c_left + Inches(0.25), Inches(4.45), c_w - Inches(0.5), Inches(2.40))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True
    p = tf_c.paragraphs[0]
    p.text = "COMPLEJIDAD Y CONDICIÓN DE NO NEGATIVIDAD"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    p = tf_c.add_paragraph()
    p.text = "• Complejidad: O((V + E) log V) implementado con min-heap binario.\n" \
             "• ¿Por qué NO admite pesos negativos?\n" \
             "    Dijkstra asume que añadir aristas no puede reducir el costo acumulado.\n" \
             "    Una arista negativa invalidaría nodos ya cerrados, produciendo rutas subóptimas.\n" \
             "    En redes físicas viales, W > 0 siempre (las distancias reales nunca son negativas)."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_LIGHT

    # Columna derecha: Imagen del asistente paso 1
    img_left = Inches(6.65)
    img_top = Inches(1.95)
    img_w = Inches(5.88)
    img_h = Inches(4.9)
    img_path = CAPTURES_ENSAYO_DIR / "02-asistente-paso1.png"
    add_image_fitted(slide, img_path, img_left, img_top, img_w, img_h,
                     caption_text="Figura 4: Paso 1 del Asistente: selección de origen en Transporte Express")

    set_speaker_notes(
        slide,
        "Para calcular la mejor ruta para cada paquete implementamos el Algoritmo de Dijkstra. Su funcionamiento se basa "
        "en una estrategia voraz o greedy: a partir de la cabecera origen, examina las aristas adyacentes y relaja las "
        "distancias tentativas hacia los vecinos. Gracias a una cola de prioridad basada en min-heap binario, su complejidad "
        "temporal es de O((V+E) log V).\n\n"
        "Es crucial enfatizar que Dijkstra exige pesos estrictamente no negativos (W > 0), lo cual se cumple naturalmente "
        "en distancias físicas de carreteras. Si existieran costos negativos o subsidios compensatorios, Dijkstra fallaría "
        "al fijar nodos permanentemente, obligando a usar algoritmos como Bellman-Ford."
    )


def build_slide_06(prs):
    """Diapositiva 6: Demo 1 (QUE -> IZA = 494.0 km)."""
    slide = create_blank_slide(prs)
    add_header(slide, "DEMOSTRACIÓN DE RUTA ÓPTIMA 1",
               "Demo 1: Quetzaltenango a Puerto Barrios (QUE → IZA)",
               "Caso de costa a costa cruzando el altiplano, región central y litoral atlántico")

    c_left = Inches(0.8)
    c_w = Inches(5.6)

    # 2 Cajas Métricas
    add_metric_box(slide, c_left, Inches(1.95), Inches(2.7), Inches(1.15),
                   "494.0 km", "Distancia Total Mínima", ACCENT_PURPLE)
    add_metric_box(slide, c_left + Inches(2.9), Inches(1.95), Inches(2.7), Inches(1.15),
                   "5 Tramos", "Carreteras Conectadas", ACCENT_GREEN)

    # Tarjeta de Secuencia de Ruta
    add_card(slide, c_left, Inches(3.25), c_w, Inches(2.20))
    tb_s = slide.shapes.add_textbox(c_left + Inches(0.2), Inches(3.35), c_w - Inches(0.4), Inches(2.0))
    tf_s = tb_s.text_frame
    tf_s.word_wrap = True
    p = tf_s.paragraphs[0]
    p.text = "SECUENCIA Y CARRETERAS OFICIALES"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    tramos = [
        ("1. QUE → TOT", "28.8 km", "RN-1 / RN-1B"),
        ("2. TOT → QUI", "36.2 km", "RD QUI-04"),
        ("3. QUI → BVE", "128.6 km", "RD QUI-2 / RN-5"),
        ("4. BVE → ZAC", "128.3 km", "CA-9 Nor / CA-14"),
        ("5. ZAC → IZA", "172.1 km", "CA-9 Nor"),
    ]
    for seg, dist, nom in tramos:
        p = tf_s.add_paragraph()
        p.text = f"• {seg}: {dist} — {nom}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_LIGHT

    # Tarjeta de Fórmula
    add_card(slide, c_left, Inches(5.60), c_w, Inches(1.35))
    tb_f = slide.shapes.add_textbox(c_left + Inches(0.2), Inches(5.70), c_w - Inches(0.4), Inches(1.15))
    tf_f = tb_f.text_frame
    tf_f.word_wrap = True
    p = tf_f.paragraphs[0]
    p.text = "FÓRMULA ARITMÉTICA DE COMPROBACIÓN"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER

    p = tf_f.add_paragraph()
    p.text = "28.8 + 36.2 + 128.6 + 128.3 + 172.1 = 494.0 km"
    p.font.name = FONT_MONO
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    # Columna derecha: Imagen
    img_left = Inches(6.65)
    img_top = Inches(1.95)
    img_w = Inches(5.88)
    img_h = Inches(4.9)
    img_path = CAPTURES_ENSAYO_DIR / "04-ejemplo-que-iza.png"
    add_image_fitted(slide, img_path, img_left, img_top, img_w, img_h,
                     caption_text="Figura 5: Ruta óptima QUE → IZA resaltada en púrpura sobre el mapa de RutaQ")

    set_speaker_notes(
        slide,
        "Veamos nuestro primer caso de demostración: un envío desde Quetzaltenango hasta Puerto Barrios, Izabal.\n\n"
        "El sistema determina una distancia óptima de 494.0 km a través de 5 tramos interdepartamentales:\n"
        "1. QUE → TOT: 28.8 km por RN-1 / RN-1B\n"
        "2. TOT → QUI: 36.2 km por RD QUI-04\n"
        "3. QUI → BVE: 128.6 km por RD QUI-2 / RN-5\n"
        "4. BVE → ZAC: 128.3 km por CA-9 Nor / CA-14\n"
        "5. ZAC → IZA: 172.1 km por CA-9 Nor\n\n"
        "La interfaz resalta la ruta sobre el mapa interactivo, desglosa cada tramo con su carretera oficial "
        "y presenta la fórmula aritmética de comprobación."
    )


def build_slide_07(prs):
    """Diapositiva 7: Demos 2 y 3 (GUA -> PET y SMA -> JUT)."""
    slide = create_blank_slide(prs)
    add_header(slide, "DEMOSTRACIONES DE RUTAS ESTRATÉGICAS 2 Y 3",
               "Demos 2 y 3: Conexiones Norte y Costa Sur",
               "Trayectos clave: acceso al departamento más extenso y corredor fronterizo")

    col_w = Inches(5.7)
    gap = Inches(0.33)

    # ------------------ Columna Izquierda: Demo 2 (GUA -> PET) ------------------
    c1_left = Inches(0.8)
    add_card(slide, c1_left, Inches(1.95), col_w, Inches(2.25))
    tb1 = slide.shapes.add_textbox(c1_left + Inches(0.2), Inches(2.05), col_w - Inches(0.4), Inches(2.05))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "DEMO 2: CIUDAD DE GUATEMALA → FLORES, PETÉN"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    p = tf1.add_paragraph()
    p.text = "• Distancia Total: 415.3 km  (3 tramos directos)\n" \
             "• Secuencia Óptima: GUA → BVE → AVE → PET\n" \
             "• Tramos Oficiales:\n" \
             "    - GUA → BVE: 89.0 km (RD BV-12 / RD GUA-12)\n" \
             "    - BVE → AVE: 69.0 km (CA-14 / RN-5)\n" \
             "    - AVE → PET: 257.3 km (RD PET-11 / RD AV-09)\n" \
             "• Fórmula: 89.0 + 69.0 + 257.3 = 415.3 km"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_LIGHT

    # Imagen Demo 2
    img1_path = CAPTURES_ENSAYO_DIR / "05-ejemplo-gua-pet.png"
    add_image_fitted(slide, img1_path, c1_left, Inches(4.30), col_w, Inches(2.55),
                     caption_text="Figura 6: Enlace capital-norte (GUA → PET = 415.3 km)")

    # ------------------ Columna Derecha: Demo 3 (SMA -> JUT) ------------------
    c2_left = c1_left + col_w + gap
    add_card(slide, c2_left, Inches(1.95), col_w, Inches(2.25))
    tb2 = slide.shapes.add_textbox(c2_left + Inches(0.2), Inches(2.05), col_w - Inches(0.4), Inches(2.05))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "DEMO 3: SAN MARCOS → JUTIAPA"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    p = tf2.add_paragraph()
    p.text = "• Distancia Total: 368.3 km  (5 tramos por Costa Sur)\n" \
             "• Secuencia Óptima: SMA → QUE → SUC → ESC → SRO → JUT\n" \
             "• Tramos Oficiales:\n" \
             "    - SMA → QUE: 49.6 km (RN-1)  |  QUE → SUC: 60.2 km (CITO-180 / SCH-02)\n" \
             "    - SUC → ESC: 102.5 km (CA-02 OCC)  |  ESC → SRO: 100.0 km (CA-2 OR / RN-16)\n" \
             "    - SRO → JUT: 56.0 km (CA-1 Or)\n" \
             "• Fórmula: 49.6 + 60.2 + 102.5 + 100.0 + 56.0 = 368.3 km"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_LIGHT

    # Imagen Demo 3
    img2_path = CAPTURES_ENSAYO_DIR / "06-ejemplo-sma-jut.png"
    add_image_fitted(slide, img2_path, c2_left, Inches(4.30), col_w, Inches(2.55),
                     caption_text="Figura 7: Corredor Costa Sur y frontera oriental (SMA → JUT = 368.3 km)")

    set_speaker_notes(
        slide,
        "En la Demo 2 resolvemos el corredor norte: Ciudad de Guatemala a Flores, Petén. La ruta óptima es "
        "de 415.3 km recorriendo tres tramos:\n"
        "- GUA → BVE: 89.0 km por RD BV-12 / RD GUA-12\n"
        "- BVE → AVE: 69.0 km por CA-14 / RN-5\n"
        "- AVE → PET: 257.3 km por RD PET-11 / RD AV-09\n\n"
        "En la Demo 3 conectamos San Marcos con Jutiapa, totalizando 368.3 km a través de cinco tramos:\n"
        "- SMA → QUE: 49.6 km por RN-1\n"
        "- QUE → SUC: 60.2 km por CITO-180 / SCH-02\n"
        "- SUC → ESC: 102.5 km por CA-02 OCC\n"
        "- ESC → SRO: 100.0 km por CA-2 OR / RN-16\n"
        "- SRO → JUT: 56.0 km por CA-1 Or\n\n"
        "Ambas rutas demuestran que el sistema encuentra de forma infalible la distancia mínima con las carreteras exactas del modelo."
    )


def build_slide_08(prs):
    """Diapositiva 8: Red Mínima (Prim y Kruskal)."""
    slide = create_blank_slide(prs)
    add_header(slide, "INFRAESTRUCTURA Y ÁRBOL DE EXPANSIÓN MÍNIMA",
               "Red Mínima Nacional: Prim y Kruskal",
               "Infraestructura vial mínima para mantener comunicados los 22 departamentos")

    c_left = Inches(0.8)
    c_w = Inches(5.6)

    # 2 Cajas Métricas
    add_metric_box(slide, c_left, Inches(1.95), Inches(2.7), Inches(1.15),
                   "21 Carreteras", "Aristas del Árbol (V - 1)", ACCENT_GREEN)
    add_metric_box(slide, c_left + Inches(2.9), Inches(1.95), Inches(2.7), Inches(1.15),
                   "1 425.3 km", "Longitud Mínima Total", ACCENT_BLUE)

    # Tarjeta de Algoritmos Prim y Kruskal
    add_card(slide, c_left, Inches(3.25), c_w, Inches(2.15))
    tb_a = slide.shapes.add_textbox(c_left + Inches(0.2), Inches(3.35), c_w - Inches(0.4), Inches(1.95))
    tf_a = tb_a.text_frame
    tf_a.word_wrap = True
    p = tf_a.paragraphs[0]
    p.text = "IMPLEMENTACIÓN DUAL DETERMINISTA"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    p = tf_a.add_paragraph()
    p.text = "• Algoritmo de Prim: Expansión voraz por corte iniciando desde una cabecera raíz, incorporando siempre la arista incidente de menor peso.\n" \
             "• Algoritmo de Kruskal: Ordenamiento global de todas las aristas con estructura Union-Find (bosque disjunto) para prevenir ciclos.\n" \
             "• Coincidencia Determinista: Ambos algoritmos producen idéntica longitud total de 1 425.3 km y las mismas 21 carreteras troncales."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_LIGHT

    # Tarjeta de Diferencia Conceptual
    add_card(slide, c_left, Inches(5.50), c_w, Inches(1.45))
    tb_d = slide.shapes.add_textbox(c_left + Inches(0.2), Inches(5.60), c_w - Inches(0.4), Inches(1.25))
    tf_d = tb_d.text_frame
    tf_d.word_wrap = True
    p = tf_d.paragraphs[0]
    p.text = "¿RUTA MÁS CORTA VS. ÁRBOL DE EXPANSIÓN MÍNIMA?"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER

    p = tf_d.add_paragraph()
    p.text = "• Ruta Más Corta (Dijkstra): Minimiza la distancia entre un par específico origen-destino.\n" \
             "• AEM (Prim/Kruskal): Minimiza la suma global de kilómetros de asfalto para conectar todos los nodos sin ciclos. Un camino dentro del AEM no siempre es la ruta más corta punto a punto."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_LIGHT

    # Columna derecha: Imagen
    img_left = Inches(6.65)
    img_top = Inches(1.95)
    img_w = Inches(5.88)
    img_h = Inches(4.9)
    img_path = CAPTURES_ENSAYO_DIR / "07-red-minima-prim.png"
    add_image_fitted(slide, img_path, img_left, img_top, img_w, img_h,
                     caption_text="Figura 8: Árbol de Expansión Mínima (1 425.3 km y 21 aristas) en el mapa nacional")

    set_speaker_notes(
        slide,
        "Además de calcular rutas para paquetes individuales, Transporte Express requería planificar su infraestructura: "
        "¿cuál es la red vial mínima que mantiene conectados los 22 departamentos con el menor kilometraje total? "
        "Esto se modela con el Árbol de Expansión Mínima (AEM).\n\n"
        "En RutaQ implementamos tanto el Algoritmo de Prim como el de Kruskal con Union-Find para evitar ciclos. "
        "Ambos algoritmos convergen con precisión determinista en una red de exactamente 21 carreteras y 1 425.3 km.\n\n"
        "La diferencia conceptual es clave: la Ruta Más Corta de Dijkstra optimiza el trayecto individual entre dos "
        "puntos, mientras que el AEM optimiza el costo global de la red completa interconectando la totalidad del territorio nacional."
    )


def build_slide_09(prs):
    """Diapositiva 9: Funcionalidades del Sistema."""
    slide = create_blank_slide(prs)
    add_header(slide, "ARQUITECTURA Y USABILIDAD DEL SOFTWARE",
               "Funcionalidades del Sistema RutaQ",
               "Herramientas especializadas para despachadores, administradores y docencia")

    # 4 Tarjetas en cuadrícula (2x2) a la izquierda y centro
    c_w = Inches(5.6)
    c_h = Inches(2.35)
    c1_left = Inches(0.8)
    c2_left = Inches(6.8)

    # 1. Asistente
    add_card(slide, c1_left, Inches(1.95), c_w, c_h)
    tb1 = slide.shapes.add_textbox(c1_left + Inches(0.2), Inches(2.05), c_w - Inches(0.4), c_h - Inches(0.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "1. ASISTENTE TRANSPORTE EXPRESS"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p = tf1.add_paragraph()
    p.text = "• Flujo guiado de 2 pasos («¿Dónde está?» y «¿A dónde va?»).\n" \
             "• Selección táctil directa sobre el mapa o mediante menús desplegables.\n" \
             "• Desglose interactivo con kilometraje acumulado, vías y fórmula de cálculo."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # 2. Administrador de Red
    add_card(slide, c2_left, Inches(1.95), c_w, c_h)
    tb2 = slide.shapes.add_textbox(c2_left + Inches(0.2), Inches(2.05), c_w - Inches(0.4), c_h - Inches(0.2))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "2. ADMINISTRADOR DE RED VIAL"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p = tf2.add_paragraph()
    p.text = "• Agregar, editar, eliminar y buscar departamentos y carreteras con validaciones.\n" \
             "• Al eliminar una carretera la ruta se recalcula al instante.\n" \
             "• Aviso y alerta automática si la red vial queda desconectada."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # 3. Tablas Paso a Paso
    add_card(slide, c1_left, Inches(4.50), c_w, c_h)
    tb3 = slide.shapes.add_textbox(c1_left + Inches(0.2), Inches(4.60), c_w - Inches(0.4), c_h - Inches(0.2))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "3. EJECUCIÓN TABULADA PASO A PASO"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_PURPLE
    p = tf3.add_paragraph()
    p.text = "• Tablas pedagógicas que replican el formato exacto enseñado en clase.\n" \
             "• Detalle de aristas consideradas, peso, nodos incorporados y descarte por ciclo.\n" \
             "• Pestañas independientes para Dijkstra, Prim, Kruskal y Floyd-Warshall."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # 4. Modo Presentación & Responsive
    add_card(slide, c2_left, Inches(4.50), c_w, c_h)
    tb4 = slide.shapes.add_textbox(c2_left + Inches(0.2), Inches(4.60), c_w - Inches(0.4), c_h - Inches(0.2))
    tf4 = tb4.text_frame
    tf4.word_wrap = True
    p = tf4.paragraphs[0]
    p.text = "4. MODO PRESENTACIÓN Y RESPONSIVE"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER
    p = tf4.add_paragraph()
    p.text = "• Vista optimizada para proyectores de aula (16:9) y pantallas widescreen.\n" \
             "• Interfaz 100 % adaptativa para teléfonos móviles, tablets y computadoras.\n" \
             "• Atajos de teclado (F para pantalla completa, flechas de navegación)."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    set_speaker_notes(
        slide,
        "RutaQ no es un prototipo estático, sino un sistema web interactivo de nivel profesional. Cuenta con cuatro "
        "módulos centrales:\n\n"
        "1. El Asistente guiado de Transporte Express en dos pasos («¿Dónde está el paquete?» y «¿A qué departamento debe entregarse?»).\n"
        "2. Un Administrador de Red para agregar, editar, eliminar y buscar departamentos y carreteras con validaciones; al eliminar una carretera la ruta se recalcula al instante y se muestra un aviso si la red queda desconectada.\n"
        "3. Tablas pedagógicas paso a paso que replican la metodología vista en clase para Prim, Kruskal, Dijkstra y Floyd-Warshall.\n"
        "4. Un Modo Presentación panorámico integrado, optimizado para proyectores de aula y pantallas táctiles móviles."
    )


def build_slide_10(prs):
    """Diapositiva 10: Verificación y Calidad."""
    slide = create_blank_slide(prs)
    add_header(slide, "AUDITORÍA TÉCNICA Y ASEGURAMIENTO DE CALIDAD",
               "Verificación Computacional y Rendimiento Web",
               "Validación científica cruzada con NetworkX y pruebas de regresión automatizadas")

    c_left = Inches(0.8)
    c_w = Inches(5.6)

    # 3 Cajas Métricas Verticales
    # NetworkX
    add_card(slide, c_left, Inches(1.95), c_w, Inches(1.50))
    tb1 = slide.shapes.add_textbox(c_left + Inches(0.2), Inches(2.05), c_w - Inches(0.4), Inches(1.30))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "1 168 / 1 168  Pares Verificados con NetworkX"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p = tf1.add_paragraph()
    p.text = "• Contraste del 100 % de caminos mínimos contra Python NetworkX: 0 discrepancias.\n" \
             "• 10 / 10 instancias de Árbol de Expansión Mínima idénticas a networkx.mst."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_LIGHT

    # Playwright
    add_card(slide, c_left, Inches(3.60), c_w, Inches(1.50))
    tb2 = slide.shapes.add_textbox(c_left + Inches(0.2), Inches(3.70), c_w - Inches(0.4), Inches(1.30))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "148 / 148  Pruebas Automatizadas (Playwright E2E)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p = tf2.add_paragraph()
    p.text = "• Suite de pruebas de regresión funcional, visual y de estado en 4 dispositivos.\n" \
             "• Pruebas ejecutadas en Desktop Chrome, iPhone 13 (WebKit), Pixel 7 e iPad."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_LIGHT

    # Lighthouse
    add_card(slide, c_left, Inches(5.25), c_w, Inches(1.65))
    tb3 = slide.shapes.add_textbox(c_left + Inches(0.2), Inches(5.35), c_w - Inches(0.4), Inches(1.45))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "Google Lighthouse: Evaluación de Rendimiento"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER
    p = tf3.add_paragraph()
    p.text = "• Rendimiento (Performance): 91 / 100\n" \
             "• Accesibilidad (Accessibility): 95 / 100\n" \
             "• Mejores Prácticas (Best Practices): 100 / 100\n" \
             "• Optimización en Motores de Búsqueda (SEO): 100 / 100"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_LIGHT

    # Columna derecha: Imagen Acerca de
    img_left = Inches(6.65)
    img_top = Inches(1.95)
    img_w = Inches(5.88)
    img_h = Inches(4.9)
    img_path = CAPTURES_ENSAYO_DIR / "10-acerca-de.png"
    add_image_fitted(slide, img_path, img_left, img_top, img_w, img_h,
                     caption_text="Figura 9: Modal Acerca de con especificaciones técnicas y créditos en RutaQ")

    set_speaker_notes(
        slide,
        "La confiabilidad de nuestro sistema está respaldada por una verificación computacional rigurosa. "
        "Contrastamos el 100 % de los caminos mínimos y árboles de expansión contra la biblioteca científica "
        "NetworkX en Python: 1 168 de 1 168 pares analizados (incluyendo los 484 pares de Guatemala y grafos de control) "
        "arrojaron idénticos resultados con cero discrepancias, y 10 de 10 instancias de AEM fueron idénticas.\n\n"
        "En aseguramiento de calidad de software, ejecutamos 148 pruebas automatizadas End-to-End con Playwright "
        "en 4 dispositivos (Desktop Chrome, iPhone 13 en WebKit, Pixel 7 e iPad) con 100 % de éxito. "
        "Finalmente, la auditoría oficial de Google Lighthouse certificó un rendimiento de 91, accesibilidad de 95, "
        "mejores prácticas de 100 y SEO de 100."
    )


def build_slide_11(prs):
    """Diapositiva 11: Conclusiones."""
    slide = create_blank_slide(prs)
    add_header(slide, "SÍNTESIS DEL PROYECTO",
               "Conclusiones y Logros Alcanzados",
               "Impacto operativo, rigor cartográfico y excelencia en ingeniería de software")

    c_w = Inches(5.6)
    c_h = Inches(2.35)
    c1_left = Inches(0.8)
    c2_left = Inches(6.8)

    # 1. Optimización Logística Real
    add_card(slide, c1_left, Inches(1.95), c_w, c_h)
    tb1 = slide.shapes.add_textbox(c1_left + Inches(0.25), Inches(2.05), c_w - Inches(0.5), c_h - Inches(0.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "1. RESOLUCIÓN DE LA PROBLEMÁTICA REAL"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p = tf1.add_paragraph()
    p.text = "• Transporte Express erradica la selección empírica de rutas por hábitos individuales.\n" \
             "• El algoritmo de Dijkstra garantiza trayectos óptimos con ahorro de combustible y reducción en tiempos de entrega.\n" \
             "• La hoja de ruta automatizada guía al conductor tramo a tramo con transparencia."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # 2. Rigor Cartográfico
    add_card(slide, c2_left, Inches(1.95), c_w, c_h)
    tb2 = slide.shapes.add_textbox(c2_left + Inches(0.25), Inches(2.05), c_w - Inches(0.5), c_h - Inches(0.2))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "2. RIGOR CARTOGRÁFICO Y REPRODUCIBILIDAD"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p = tf2.add_paragraph()
    p.text = "• Grafo no dirigido de 22 cabeceras y 49 carreteras directas sustentado en OpenStreetMap.\n" \
             "• Auditoría 49/49 en Google Maps que valida la precisión de cada distancia utilizada.\n" \
             "• Inclusión estratégica del tramo IZA–PET (287.0 km) por la CA-13 para evitar desvíos falsos."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # 3. Sinergia Algorítmica
    add_card(slide, c1_left, Inches(4.50), c_w, c_h)
    tb3 = slide.shapes.add_textbox(c1_left + Inches(0.25), Inches(4.60), c_w - Inches(0.5), c_h - Inches(0.2))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "3. DOMINIO ALGORÍTMICO INTEGRAL"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_PURPLE
    p = tf3.add_paragraph()
    p.text = "• Coexistencia armónica: Dijkstra para encomiendas puntuales y Prim/Kruskal para red troncal.\n" \
             "• Ambos algoritmos de MST coinciden deterministamente en 21 carreteras y 1 425.3 km.\n" \
             "• Comprensión clara de la frontera teórica: por qué Dijkstra no admite pesos negativos."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # 4. Excelencia en Software
    add_card(slide, c2_left, Inches(4.50), c_w, c_h)
    tb4 = slide.shapes.add_textbox(c2_left + Inches(0.25), Inches(4.60), c_w - Inches(0.5), c_h - Inches(0.2))
    tf4 = tb4.text_frame
    tf4.word_wrap = True
    p = tf4.paragraphs[0]
    p.text = "4. CALIDAD DE INGENIERÍA DE SOFTWARE"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER
    p = tf4.add_paragraph()
    p.text = "• Arquitectura modular 100 % TypeScript, reactiva y ejecutable directamente en el navegador.\n" \
             "• Verificación cruzada contra NetworkX (1 168/1 168) y 148 pruebas automatizadas Playwright.\n" \
             "• Despliegue web activo, accesible y con código abierto para auditoría comunitaria."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    set_speaker_notes(
        slide,
        "En conclusión, el proyecto demuestra el poder de la Investigación de Operaciones aplicada a la logística real:\n\n"
        "1. Transporte Express erradica la improvisación de rutas sustituyendo la intuición por algoritmos deterministas.\n"
        "2. El modelo matemático está fundamentado en datos reales de los 22 departamentos y 49 carreteras auditadas.\n"
        "3. Se integraron sinérgicamente Dijkstra para el despacho diario y Prim/Kruskal para la planeación de la red troncal (1 425.3 km).\n"
        "4. El código en TypeScript es modular, robusto y validado contra bibliotecas científicas de la industria, "
        "cumpliendo con creces los requerimientos del curso."
    )


def build_slide_12(prs):
    """Diapositiva 12: Gracias y Preguntas."""
    slide = create_blank_slide(prs)

    # Tarjeta Principal Centrada
    card_w = Inches(11.733)
    card_h = Inches(6.3)
    card_left = Inches(0.8)
    card_top = Inches(0.6)
    add_card(slide, card_left, card_top, card_w, card_h, bg_color=BG_CARD, border_color=BORDER_ACCENT)

    # Insignia
    add_badge(slide, "SESIÓN DE PREGUNTAS Y DEMOSTRACIÓN EN VIVO",
              card_left + Inches(0.6), card_top + Inches(0.50),
              width=Inches(5.2), height=Inches(0.34),
              fg_color=ACCENT_GREEN, bg_color=BADGE_BG_GREEN)

    # Título
    title_box = slide.shapes.add_textbox(card_left + Inches(0.6), card_top + Inches(0.95), Inches(10.5), Inches(0.9))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "¡Muchas Gracias por su Atención!"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(34)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    p_sub = tf.add_paragraph()
    p_sub.text = "Transporte Express — Sistema de Optimización de Rutas (RutaQ)"
    p_sub.font.name = FONT_FAMILY
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = ACCENT_CYAN

    # Tarjeta de Enlaces Oficiales (Izquierda)
    link_card = add_card(slide, card_left + Inches(0.6), card_top + Inches(2.10), Inches(5.2), Inches(3.40),
                         bg_color=BG_CARD_ALT, border_color=BORDER_CARD)
    tb_l = slide.shapes.add_textbox(card_left + Inches(0.8), card_top + Inches(2.25), Inches(4.8), Inches(3.10))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "ACCESO AL SISTEMA Y DOCUMENTACIÓN"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER

    links_info = [
        ("Aplicación en Vivo:", "https://criss-paz.github.io/rutaq/"),
        ("Presentación Web:", "https://criss-paz.github.io/rutaq/presentacion/"),
        ("Repositorio GitHub:", "https://github.com/Criss-paz/rutaq"),
        ("Catedrático:", "Ing. Samuel Isaac Gómez Moreno"),
        ("Curso:", "Investigación de Operaciones · UMG 2026"),
    ]
    for lbl, val in links_info:
        p = tf_l.add_paragraph()
        p.text = f"{lbl}\n{val}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_LIGHT

    # Tarjeta de Demostración y Preguntas (Derecha)
    add_card(slide, card_left + Inches(6.1), card_top + Inches(2.10), Inches(5.0), Inches(3.40),
             bg_color=BG_CARD_ALT, border_color=ACCENT_PURPLE)

    img_path = CAPTURES_DIR / "11-modo-presentacion.png"
    add_image_fitted(slide, img_path, card_left + Inches(6.3), card_top + Inches(2.25), Inches(4.6), Inches(2.3),
                     caption_text="RutaQ en Modo Presentación listo para pruebas en vivo")

    tb_q = slide.shapes.add_textbox(card_left + Inches(6.3), card_top + Inches(4.90), Inches(4.6), Inches(0.55))
    tf_q = tb_q.text_frame
    p_q = tf_q.paragraphs[0]
    p_q.text = "¿Alguna pregunta o demostración en vivo adicional?"
    p_q.font.name = FONT_FAMILY
    p_q.font.size = Pt(13)
    p_q.font.bold = True
    p_q.font.color.rgb = ACCENT_PURPLE
    p_q.alignment = PP_ALIGN.CENTER

    # Pie
    footer_box = slide.shapes.add_textbox(card_left + Inches(0.6), card_top + Inches(5.70), Inches(10.5), Inches(0.4))
    tf_ft = footer_box.text_frame
    p_ft = tf_ft.paragraphs[0]
    p_ft.text = "Cristopher Castellanos (2690-24-5972) · Leonel Taracena (2690-24-4072) · Pablo Vásquez (2690-24-13123)"
    p_ft.font.name = FONT_FAMILY
    p_ft.font.size = Pt(11)
    p_ft.font.color.rgb = TEXT_MUTED
    p_ft.alignment = PP_ALIGN.CENTER

    set_speaker_notes(
        slide,
        "Muchas gracias catedrático y compañeros por su tiempo y atención. El sistema se encuentra completamente "
        "desplegado y operativo en la web para que cualquiera pueda probarlo en su teléfono o computadora en "
        "criss-paz.github.io/rutaq/, y la versión web de estas diapositivas está en /presentacion/.\n\n"
        "El código fuente completo y los reportes de verificación están abiertos en GitHub. "
        "Quedamos a su entera disposición para responder cualquier pregunta o realizar las pruebas de ruta "
        "en vivo que deseen solicitar."
    )


# -----------------------------------------------------------------------------
# Función Principal
# -----------------------------------------------------------------------------

def main():
    ensure_dirs()
    print("=" * 70)
    print("GENERADOR DE PRESENTACIÓN — TRANSPORTE EXPRESS (RutaQ)")
    print("=" * 70)
    print(f"Directorio de salida: {OUTPUT_DIR}")

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH_IN)
    prs.slide_height = Inches(SLIDE_HEIGHT_IN)

    builders = [
        ("Diapositiva 01: Portada", build_slide_01),
        ("Diapositiva 02: El Problema Logístico", build_slide_02),
        ("Diapositiva 03: Modelo Matemático del Grafo", build_slide_03),
        ("Diapositiva 04: Datos Reales y Validación", build_slide_04),
        ("Diapositiva 05: Algoritmo de Dijkstra", build_slide_05),
        ("Diapositiva 06: Demo 1 (QUE → IZA = 494.0 km)", build_slide_06),
        ("Diapositiva 07: Demos 2 y 3 (GUA → PET y SMA → JUT)", build_slide_07),
        ("Diapositiva 08: Red Mínima Nacional (Prim y Kruskal)", build_slide_08),
        ("Diapositiva 09: Funcionalidades del Sistema", build_slide_09),
        ("Diapositiva 10: Verificación y Calidad", build_slide_10),
        ("Diapositiva 11: Conclusiones", build_slide_11),
        ("Diapositiva 12: Gracias y Preguntas", build_slide_12),
    ]

    for title, builder in builders:
        print(f"  [+] Generando {title}...")
        builder(prs)

    prs.save(str(OUTPUT_FILE))
    print("-" * 70)
    print(f"[ÉXITO] Presentación guardada exitosamente en:\n  {OUTPUT_FILE}")
    print(f"Total diapositivas: {len(prs.slides)}")
    print("=" * 70)


if __name__ == "__main__":
    main()
